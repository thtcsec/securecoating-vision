"""Surgical Application Form identity and claim fixes."""

from __future__ import annotations

from pathlib import Path

import yaml
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Al + Materials Competition Application Form.docx"
IDENTITY = ROOT / "configs" / "project_identity.yaml"

# Explicit competition revision date for this organizer-approved title sync (DD / MM / YYYY).
SUBMISSION_REVISION_DATE = "07 / 09 / 2026"

PERFORMANCE_PREDICTION_BOUNDARY = (
    "SecureCoating covers the inspection-to-quality-decision segment of an autonomous "
    "materials characterization pipeline; material/electrochemical performance prediction "
    "is the next validation stage once plant-linked property labels exist."
)
FACTORY_NONCLAIM = (
    "The repository does not claim factory qualification or production performance."
)
PERFORMANCE_BOUNDARY_WITH_NONCLAIM = (
    f"{PERFORMANCE_PREDICTION_BOUNDARY} {FACTORY_NONCLAIM}"
)

# Idempotent string rewrites: each `old` must NOT remain as a substring of `new`.
REPLACEMENTS = (
    (
        "certificates use canonical HMAC-SHA256 payloads",
        "certificates use HMAC-SHA256 authenticated traceability tags",
    ),
    (
        "and can be represented in a signed certificate snapshot.",
        "and can be represented in an HMAC-SHA256 authenticated certificate snapshot.",
    ),
    (
        "This is a validation extension, not a change of topic.",
        "The LIBAD lane extends validation of the same battery-electrode inspection and quality-decision problem.",
    ),
    (
        "Next Validation Steps: Obtain an independent roll-disjoint dataset, hash-verify "
        "official LIBAD, complete PLC/HIL and safety testing, validate factory calibration, "
        "and publish only measurements supported by immutable evidence.",
        "Next Validation Steps: Obtain an independent roll-disjoint factory dataset; complete "
        "PLC/HIL, plant calibration, and safety testing; pursue authors' DINOv3/DA-Core "
        "validation only if paper-comparable LIBAD benchmarking is required; and treat "
        "downstream material/electrochemical performance prediction as future validation, "
        "not a current defense claim.",
    ),
    (
        "Next Validation Steps: Reproduce the official-LIBAD validation under the current "
        "clean release and preserve current-source provenance; obtain an independent "
        "roll-disjoint factory dataset; complete PLC/HIL and plant calibration; and keep "
        "any future performance-risk proxy explicitly simulation-validated and separate "
        "from electrochemical cell-performance claims.",
        "Next Validation Steps: Obtain an independent roll-disjoint factory dataset; complete "
        "PLC/HIL, plant calibration, and safety testing; pursue authors' DINOv3/DA-Core "
        "validation only if paper-comparable LIBAD benchmarking is required; and treat "
        "downstream material/electrochemical performance prediction as future validation, "
        "not a current defense claim.",
    ),
)

SECTION7_HEADING = "7、Team Member Information"

ADVISOR_LINE = (
    "Advisor / 指导老师: Prof. Kris Singh — Visiting Professor, Tsinghua University; "
    "Founder & CEO, SRII. Advisory scope: innovation, industrialization, and "
    "final-defense guidance."
)

CONTESTANT_LEADER = {
    "role": "Team Leader",
    "full_name": "Trịnh Hoàng Tú",
    "affiliation": "HUFLIT",
    "phone": "+84 983967098",
    "email": "tht.csec2005@gmail.com",
}

# Nested roster must fit the usable page content width (~14.7 cm). Auto-layout was
# squeezing the email column and wrapping the final "m" onto a new line.
SECTION7_COLUMN_WIDTHS = (
    Cm(1.9),  # Role
    Cm(2.9),  # Full Name
    Cm(2.3),  # Affiliation
    Cm(2.5),  # Phone
    Cm(4.2),  # Email — keeps tht.csec2005@gmail.com on one line at 8 pt
)
EMAIL_CELL_FONT_PT = 8.0


def ensure_performance_prediction_boundary(text: str) -> tuple[str, int]:
    """Ensure the performance-prediction boundary sentence appears exactly once.

    Idempotent: repeated application neither inserts extra copies nor changes an
    already-normalized paragraph that already contains the combined target.
    """
    changes = 0
    updated = text

    # Collapse consecutive duplicates first (repairs prior non-idempotent runs).
    doubled = f"{PERFORMANCE_PREDICTION_BOUNDARY} {PERFORMANCE_PREDICTION_BOUNDARY}"
    while doubled in updated:
        updated = updated.replace(doubled, PERFORMANCE_PREDICTION_BOUNDARY)
        changes += 1

    if PERFORMANCE_BOUNDARY_WITH_NONCLAIM in updated:
        # Already in the desired state (boundary immediately precedes non-claim).
        return updated, changes

    if FACTORY_NONCLAIM in updated:
        # Insert only when the boundary does not already precede the non-claim.
        updated = updated.replace(FACTORY_NONCLAIM, PERFORMANCE_BOUNDARY_WITH_NONCLAIM, 1)
        changes += 1
        while doubled in updated:
            updated = updated.replace(doubled, PERFORMANCE_PREDICTION_BOUNDARY)
            changes += 1

    return updated, changes


def _set_cell_text(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        return
    for paragraph in list(paragraphs[1:]):
        parent = paragraph._element.getparent()
        if parent is not None:
            parent.remove(paragraph._element)
    paragraphs[0].clear()
    lines = text.split("\n")
    paragraphs[0].add_run(lines[0] if lines else "")
    for line in lines[1:]:
        cell.add_paragraph(line)


def _strip_advisor_block(text: str) -> str:
    keep = []
    for line in text.splitlines():
        s = line.strip()
        if (
            s.startswith("Advisor")
            or s.startswith("指导老师")
            or s.startswith("Advisory scope")
            or "kris@thesrii.org" in s
        ):
            continue
        keep.append(line)
    return "\n".join(keep).rstrip()


def _set_nested_cell(cell, text: str, *, font_pt: float | None = None) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.text = text
        return
    for paragraph in list(paragraphs[1:]):
        parent = paragraph._element.getparent()
        if parent is not None:
            parent.remove(paragraph._element)
    paragraphs[0].clear()
    run = paragraphs[0].add_run(text)
    if font_pt is not None:
        run.font.size = Pt(font_pt)


def _dxa(width) -> str:
    """Word dxa (twips) string from a python-docx Length/Cm value."""
    return str(int(width) // 635)


def _set_cell_width(cell, width) -> None:
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), _dxa(width))
    tc_w.set(qn("w:type"), "dxa")


def _clear_paragraph_text(paragraph, replacement: str | None = None) -> None:
    for run in paragraph.runs:
        run.text = ""
    if replacement is None:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replacement
    else:
        paragraph.add_run(replacement)


def _fix_section7_contestant_roster(cell) -> dict:
    """Keep only the contestant roster in the nested table; advisor stays in Section 6.

    Important: do not delete paragraph XML nodes in this cell. Section 7 embeds a nested
    table; removing sibling paragraphs via python-docx corrupts the OOXML for Microsoft Word.
    """
    stats = {
        "heading_set": False,
        "nested_rows_cleared": 0,
        "leader_set": False,
        "plain_roster_cleared": 0,
        "email_layout_fixed": False,
    }
    heading_set = False
    for paragraph in cell.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        if text.startswith("7、") and not heading_set:
            _clear_paragraph_text(paragraph, SECTION7_HEADING)
            heading_set = True
            stats["heading_set"] = True
            continue
        if text.startswith("Team Leader") or text.startswith("Team Member"):
            _clear_paragraph_text(paragraph, "")
            stats["plain_roster_cleared"] += 1

    if not cell.tables:
        raise SystemExit("Section 7 nested roster table missing")

    roster = cell.tables[0]
    if len(roster.rows) < 2 or len(roster.columns) < 5:
        raise SystemExit("Section 7 nested roster table has unexpected shape")

    # Fixed layout inside usable page width so Word cannot squeeze the email column.
    # Note: Cm+Cm returns a raw EMU int in python-docx, so sum via int().
    total_emu = sum(int(width) for width in SECTION7_COLUMN_WIDTHS)
    roster.autofit = False
    tbl_pr = roster._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        roster._tbl.insert(0, tbl_pr)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_emu // 635))
    tbl_w.set(qn("w:type"), "dxa")
    for row in roster.rows:
        for idx, width in enumerate(SECTION7_COLUMN_WIDTHS):
            if idx < len(row.cells):
                _set_cell_width(row.cells[idx], width)
    stats["email_layout_fixed"] = True

    leader = roster.rows[1]
    _set_nested_cell(leader.cells[0], CONTESTANT_LEADER["role"])
    _set_nested_cell(leader.cells[1], CONTESTANT_LEADER["full_name"])
    _set_nested_cell(leader.cells[2], CONTESTANT_LEADER["affiliation"])
    _set_nested_cell(leader.cells[3], CONTESTANT_LEADER["phone"])
    _set_nested_cell(
        leader.cells[4],
        CONTESTANT_LEADER["email"],
        font_pt=EMAIL_CELL_FONT_PT,
    )
    stats["leader_set"] = True

    for row in list(roster.rows)[2:]:
        name = (row.cells[1].text or "").strip().lower()
        role = (row.cells[0].text or "").strip().lower()
        if "kris" in name or "advisor" in name or role == "team member" or name:
            for cell_i in range(5):
                _set_nested_cell(row.cells[cell_i], "")
            _set_nested_cell(row.cells[0], "Team Member")
            stats["nested_rows_cleared"] += 1

    roster_blob = "\n".join(c.text for r in roster.rows for c in r.cells).lower()
    if "kris" in roster_blob or "advisor" in roster_blob or "srii" in roster_blob:
        raise SystemExit("refusing to save: advisor identity still present in Section 7 roster")
    return stats


def _set_revision_date(doc: Document, date_text: str = SUBMISSION_REVISION_DATE) -> bool:
    updated = False
    for paragraph in doc.paragraphs:
        text = paragraph.text or ""
        if text.strip().startswith("Date:"):
            target = f"Date: {date_text}"
            if text.strip() != target:
                _clear_paragraph_text(paragraph, target)
                updated = True
            break
    return updated


def fix(path: Path = DOCX) -> dict:
    identity = yaml.safe_load(IDENTITY.read_text(encoding="utf-8"))
    new_title = str(identity["registered_title"]).strip()
    doc = Document(str(path))
    if len(doc.tables) < 2:
        raise SystemExit("expected cover + body tables")

    stats = {
        "title_updated": False,
        "replacements": 0,
        "boundary_normalizations": 0,
        "section6": False,
        "section7": False,
        "date_updated": False,
        "section7_roster": {},
    }

    cover = doc.tables[0].rows[0].cells[1]
    if cover.text.strip() != new_title:
        _set_cell_text(cover, new_title)
        stats["title_updated"] = True

    # Cover: Team Member's Name(s) remains N/A (no contestant teammates).
    member_cell = doc.tables[0].rows[3].cells[1]
    if "kris" in member_cell.text.lower() or "singh" in member_cell.text.lower():
        _set_cell_text(member_cell, "N/A")

    body = doc.tables[1]
    for row in body.rows:
        cell = row.cells[0]
        text = cell.text or ""
        updated = text
        for old, new in REPLACEMENTS:
            if old in updated:
                updated = updated.replace(old, new)
                stats["replacements"] += 1

        updated, boundary_changes = ensure_performance_prediction_boundary(updated)
        stats["boundary_normalizations"] += boundary_changes

        if updated.strip().startswith("6、Additional Information"):
            updated = _strip_advisor_block(updated)
            if ADVISOR_LINE not in updated:
                updated = updated.rstrip() + "\n\n" + ADVISOR_LINE
            _set_cell_text(cell, updated)
            stats["section6"] = True
            continue

        if updated.strip().startswith("7、Team Member Information"):
            stats["section7_roster"] = _fix_section7_contestant_roster(cell)
            stats["section7"] = True
            continue

        if updated != text:
            _set_cell_text(cell, updated)

    stats["date_updated"] = _set_revision_date(doc)

    title = doc.tables[0].rows[0].cells[1].text.strip()
    if title != new_title:
        raise SystemExit(f"refusing to save: cover title mismatch: {title!r}")
    if "High-Throughput and Zero-Trust" in title:
        raise SystemExit("refusing to save: old competition title still present")

    # Final semantic guards.
    full = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    boundary_count = full.count(PERFORMANCE_PREDICTION_BOUNDARY)
    if boundary_count != 1:
        raise SystemExit(
            f"refusing to save: performance-prediction boundary must appear exactly once "
            f"(found {boundary_count})"
        )
    if PERFORMANCE_BOUNDARY_WITH_NONCLAIM not in full:
        raise SystemExit(
            "refusing to save: performance-prediction boundary must precede factory non-claim"
        )
    section7 = doc.tables[1].rows[6].cells[0]
    roster_text = ""
    if section7.tables:
        roster_text = "\n".join(c.text for r in section7.tables[0].rows for c in r.cells)
    if "Kris" in roster_text or "kris@thesrii.org" in roster_text:
        raise SystemExit("refusing to save: Kris Singh appears in Section 7 contestant roster")
    if "Team Leader —" in section7.text or "Team Member —" in section7.text:
        raise SystemExit("refusing to save: redundant plain-text roster lines remain in Section 7")
    if ADVISOR_LINE.split(":")[0] not in full and "Prof. Kris Singh" not in full:
        raise SystemExit("refusing to save: advisor line missing from Application Form")
    email = ""
    if section7.tables:
        email = section7.tables[0].rows[1].cells[4].text.strip()
    if email != CONTESTANT_LEADER["email"]:
        raise SystemExit(f"refusing to save: Section 7 email mismatch: {email!r}")

    # Scrub machine/vendor core properties for provenance-facing submission packs.
    cp = doc.core_properties
    cp.author = "Trinh Hoang Tu"
    cp.last_modified_by = "Trinh Hoang Tu"
    cp.title = "AI + Materials Competition Application Form — Team 71"
    cp.subject = "Track 4 submission — SecureCoating-Vision"

    doc.save(str(path))
    return stats


def main() -> None:
    import json

    print(json.dumps(fix(), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
