"""Generate the 6-minute final-defense deck from the registered identity.

Clean typography title slide (Team 71 / HUFLIT). No organizer lab logo as
product branding. Honest current evidence, not retired leftover decks.

Optional extra (not an API runtime dependency):
    .venv\\Scripts\\python.exe -m pip install python-pptx
    .venv\\Scripts\\python.exe scripts/generate_defense_slides.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "SecureCoating-Vision_Final_Defense_6min.pptx"
EXTERNAL_DEMO = ROOT / "reports/external_coatingvision_demo/coatingvision_model_output.png"
RGB_GIF = ROOT / "reports/defense_gifs/rgb_hold_replay.gif"
LIBAD_GIF = ROOT / "reports/defense_gifs/libad_gate.gif"
HELD_OUT_SURFACE = ROOT / "reports/defense_gifs/coating_surface_heldout.png"
RGB_METRICS = ROOT / "reports/coatingvision_real_test_metrics.json"

BG = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF4, 0xF7, 0xFB)
INK = RGBColor(0x1A, 0x1F, 0x2B)
ACCENT = RGBColor(0x1F, 0x5E, 0xB8)
GREEN = RGBColor(0x0F, 0x7B, 0x45)
HOLD = RGBColor(0xC4, 0x7B, 0x00)
RED = RGBColor(0xC8, 0x10, 0x2E)
MUTED = RGBColor(0x5C, 0x67, 0x75)
LINE = RGBColor(0xD5, 0xDC, 0xE6)
RULE = RGBColor(0x1F, 0x5E, 0xB8)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TITLE = (
    "SecureCoating Vision: A High-Throughput and Zero-Trust Edge-Cloud "
    "Pipeline for Inline Battery Electrode Defect Inspection and Traceable "
    "Quality Decisions"
)
TAGLINE = "Evidence-Gated Multimodal Inspection for Battery Electrode Manufacturing"


def _load_rgb_metrics() -> dict:
    import json

    payload = json.loads(RGB_METRICS.read_text(encoding="utf-8"))
    metrics = payload["metrics"]
    provenance = payload.get("dataset_provenance") or {}
    weights = str(payload.get("weights_sha256") or "")
    dataset = str(provenance.get("dataset_tree_sha256") or "")
    if len(weights) < 8 or len(dataset) < 8:
        raise ValueError("coatingvision_real_test_metrics.json missing weight/dataset hashes")
    return {
        "precision_pct": f"{float(metrics['metrics/precision(B)']) * 100:.1f}%",
        "recall_pct": f"{float(metrics['metrics/recall(B)']) * 100:.1f}%",
        "map50_pct": f"{float(metrics['metrics/mAP50(B)']) * 100:.1f}%",
        "map50_95_pct": f"{float(metrics['metrics/mAP50-95(B)']) * 100:.1f}%",
        "map50_short": f"{float(metrics['metrics/mAP50(B)']):.3f}",
        "n_test": int((provenance.get("splits") or {}).get("test") or 88),
        "seed": provenance.get("seed", 71),
        "weights_prefix": weights[:8],
        "dataset_prefix": dataset[:8],
    }


def _picture_size_px(path: Path) -> tuple[int, int]:
    from PIL import Image

    with Image.open(path) as image:
        return int(image.size[0]), int(image.size[1])


def _length_inches(value) -> float:
    inches = getattr(value, "inches", None)
    if inches is not None:
        return float(inches)
    return float(value)


def _add_picture_contain(slide, path: Path, left, top, max_width, max_height):
    """Fit media inside a box without stretching or overflowing the CLOSE band."""
    px_w, px_h = _picture_size_px(path)
    aspect = px_w / float(px_h)
    box_w = _length_inches(max_width)
    box_h = _length_inches(max_height)
    if box_w / box_h > aspect:
        height = box_h
        width = box_h * aspect
    else:
        width = box_w
        height = box_w / aspect
    # Top-align inside the media frame so CLOSE text below stays clear.
    x = _length_inches(left) + (box_w - width) / 2.0
    y = _length_inches(top)
    slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), width=Inches(width), height=Inches(height)
    )


def _set_run(run, text, size_pt, color, bold=False):
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _box(slide, left, top, width, height, fill=CARD, accent=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill(shape, fill)
    shape.line.color.rgb = LINE
    shape.line.width = Emu(6350)
    if accent:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
        _fill(bar, ACCENT)
    return shape


def _rule(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Emu(12700))
    _fill(shape, RULE)
    return shape


def add_textbox(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(item.get("space_after", 6))
        run = paragraph.add_run()
        _set_run(run, item["text"], item.get("size", 16), item.get("color", INK), item.get("bold", False))
    return box


def footer(slide, number):
    add_textbox(
        slide,
        Inches(0.45),
        Inches(7.12),
        Inches(8.5),
        Inches(0.28),
        [{"text": "Team 71  ·  Trịnh Hoàng Tú  ·  HUFLIT  ·  Track 4", "size": 11, "color": MUTED}],
    )
    add_textbox(
        slide,
        Inches(10.6),
        Inches(7.12),
        Inches(2.2),
        Inches(0.28),
        [{"text": f"TEAM 71  /  {number}", "size": 11, "color": MUTED, "bold": True}],
        align=PP_ALIGN.RIGHT,
    )


def paint_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    _fill(shape, BG)
    bottom = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.05), SLIDE_W, Emu(12700))
    _fill(bottom, LINE)


def kicker(slide, code, timing):
    add_textbox(
        slide,
        Inches(0.45),
        Inches(0.22),
        Inches(12.4),
        Inches(0.32),
        [{"text": f"{code}    |    {timing}", "size": 12, "color": ACCENT, "bold": True}],
    )


def build() -> Path:
    rgb = _load_rgb_metrics()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    add_textbox(s, Inches(0.45), Inches(0.35), Inches(12.4), Inches(0.35), [
        {
            "text": "Global AI + Materials Innovation Application Competition 2026",
            "size": 14,
            "color": ACCENT,
            "bold": True,
        }
    ])
    add_textbox(s, Inches(0.45), Inches(1.05), Inches(12.4), Inches(0.45), [
        {"text": "SecureCoating Vision", "size": 32, "color": INK, "bold": True}
    ])
    add_textbox(s, Inches(0.45), Inches(1.55), Inches(12.4), Inches(1.35), [
        {"text": TITLE, "size": 17, "color": INK, "bold": True, "space_after": 0}
    ])
    add_textbox(s, Inches(0.45), Inches(3.05), Inches(12.4), Inches(0.4), [
        {"text": TAGLINE, "size": 18, "color": ACCENT, "bold": True}
    ])
    _rule(s, Inches(0.45), Inches(3.55), Inches(12.4))
    add_textbox(s, Inches(0.45), Inches(3.85), Inches(12.4), Inches(0.9), [
        {
            "text": "Team 71 · Track 4 — AI + Materials Testing & Characterization",
            "size": 18,
            "color": INK,
            "bold": True,
            "space_after": 8,
        },
        {
            "text": "Trinh Hoang Tu · HUFLIT",
            "size": 18,
            "color": INK,
            "bold": True,
            "space_after": 10,
        },
        {
            "text": (
                "Advisor: Prof. Kris Singh · Visiting Professor, Tsinghua University "
                "· Founder & CEO, SRII"
            ),
            "size": 13,
            "color": MUTED,
            "bold": False,
        },
    ])
    add_textbox(s, Inches(0.45), Inches(5.55), Inches(12.4), Inches(1.1), [
        {
            "text": (
                "Registered title describes the target Zero-Trust Edge-Cloud architecture. "
                "This prototype validates inspection, evidence gating, and fail-closed contracts — "
                "not completed mTLS/RBAC/OT or factory qualification."
            ),
            "size": 14,
            "color": MUTED,
            "space_after": 8,
        },
        {
            "text": "Decision objective: prove a defensible inline inspection architecture. Supervised factory pilot is the next gate.",
            "size": 14,
            "color": MUTED,
        },
    ])
    footer(s, 1)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "02", "0:35–1:15")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(7.4), Inches(1.1), [
        {"text": "The defect is tiny. An automatic PASS is not.", "size": 28, "color": INK, "bold": True}
    ])
    add_textbox(s, Inches(0.45), Inches(1.7), Inches(7.4), Inches(2.4), [
        {"text": "Electrode coating defects become scrap, rework, or untraceable cell risk.", "size": 18, "color": INK, "space_after": 10},
        {"text": "A detector that ranks anomalies well can still be unsafe to act on if sensors disagree, calibration is unverified, or the PLC does not acknowledge.", "size": 16, "color": MUTED, "space_after": 10},
        {"text": "Academic detection answers how to detect. A factory needs when a detection is safe enough to act on.", "size": 16, "color": ACCENT, "bold": True},
    ])
    _box(s, Inches(0.45), Inches(4.3), Inches(7.4), Inches(2.4), accent=True)
    add_textbox(s, Inches(0.65), Inches(4.45), Inches(7.0), Inches(2.1), [
        {"text": "MEASURED ON REAL OPTICAL DATA", "size": 12, "color": RED, "bold": True},
        {"text": f"{rgb['n_test']} held-out images · mAP50 {rgb['map50_short']}", "size": 18, "color": INK, "bold": True, "space_after": 8},
        {"text": f"Public CoatingVision image-disjoint split, seed {rgb['seed']}. Not factory roll-disjoint; HIL and calibration remain pilot gates.", "size": 14, "color": MUTED},
    ])
    if HELD_OUT_SURFACE.is_file():
        _box(s, Inches(8.15), Inches(0.55), Inches(4.7), Inches(6.2))
        s.shapes.add_picture(str(HELD_OUT_SURFACE), Inches(8.3), Inches(0.95), width=Inches(4.4))
        add_textbox(s, Inches(8.3), Inches(5.85), Inches(4.4), Inches(0.7), [
            {"text": "Real coating surface (held-out CoatingVision sample). RGB-only. Not a multimodal plant capture.", "size": 12, "color": MUTED}
        ])
    footer(s, 2)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "03", "1:15–2:00")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(12.4), Inches(0.7), [
        {"text": "Two evidence lanes. One fail-closed contract.", "size": 28, "color": INK, "bold": True}
    ])
    _box(s, Inches(0.45), Inches(1.4), Inches(6.05), Inches(5.3), accent=True)
    add_textbox(s, Inches(0.65), Inches(1.55), Inches(5.65), Inches(4.9), [
        {"text": "LANE A  ·  IMPLEMENTED", "size": 12, "color": GREEN, "bold": True},
        {"text": "RGB YOLO26n detect / ONNX", "size": 22, "color": INK, "bold": True},
        {"text": "FastAPI  ·  SQLite PENDING→final  ·  OPC UA / Modbus command+ACK  ·  HMAC certificate", "size": 15, "color": MUTED, "space_after": 12},
        {"text": "Thermal and profilometry are simulated adapters. They are not plant-instrument measurements.", "size": 15, "color": HOLD},
    ])
    _box(s, Inches(6.8), Inches(1.4), Inches(6.05), Inches(5.3), accent=True)
    add_textbox(s, Inches(7.0), Inches(1.55), Inches(5.65), Inches(4.9), [
        {"text": "LANE B  ·  VALIDATION EXTENSION", "size": 12, "color": ACCENT, "bold": True},
        {"text": "LIBAD VIS + X-rayL adapter", "size": 22, "color": INK, "bold": True},
        {"text": "Official mount hashed. Numpy 10-seed ~AUROC 0.70 / FPR95 0.84. Authors' runner interim DINOv2 (1 seed) ~AUROC 0.856 / FPR95 0.716 — still not paper DINOv3. Demo cases stay protocol fixtures.", "size": 15, "color": MUTED, "space_after": 12},
        {"text": "DA-Core is Sui et al. SecureCoating-Vision adds the evidence gate, not a new detector claim.", "size": 15, "color": ACCENT},
    ])
    footer(s, 3)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "04", "2:00–2:45")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(12.4), Inches(0.7), [
        {"text": "Model output cannot self-release the line.", "size": 28, "color": INK, "bold": True}
    ])
    for i, (name, color, body) in enumerate((
        ("PASS", GREEN, "OPTIMAL inference, trained model, verified calibration, healthy traceability, PLC ACK."),
        ("REJECT", RED, "Evidence supports a defect and the communication contract completes."),
        ("HOLD", HOLD, "Disagreement, stale/missing sensor, timeout, unverified calibration, DB fault, or unconfirmed PLC."),
    )):
        left = Inches(0.45 + i * 4.2)
        _box(s, left, Inches(1.4), Inches(4.0), Inches(2.4), accent=True)
        add_textbox(s, left + Inches(0.2), Inches(1.55), Inches(3.6), Inches(0.45), [
            {"text": name, "size": 26, "color": color, "bold": True}
        ])
        add_textbox(s, left + Inches(0.2), Inches(2.15), Inches(3.6), Inches(1.4), [
            {"text": body, "size": 15, "color": INK}
        ])
    _box(s, Inches(0.45), Inches(4.05), Inches(12.4), Inches(2.65), accent=True)
    add_textbox(s, Inches(0.7), Inches(4.2), Inches(12.0), Inches(2.3), [
        {"text": "SOFTWARE SAFETY SEMANTICS", "size": 12, "color": ACCENT, "bold": True},
        {"text": "Uncertainty becomes a controlled software state — not a claim of factory-qualified process safety.", "size": 18, "color": INK, "bold": True, "space_after": 10},
        {"text": "Readiness failure → HOLD. PASS/REJECT require confirmed control evidence. No readiness gate, no automatic release. Software E-stop latches local interlock and requests PLC channels; it is not a safety-rated hardwired stop.", "size": 14, "color": MUTED},
    ])
    footer(s, 4)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "05", "2:45–3:35")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(12.4), Inches(0.7), [
        {"text": "Readiness is visible before any line decision.", "size": 26, "color": INK, "bold": True}
    ])
    for i, (title, body) in enumerate((
        ("OPERATE", "Line disposition, model hashes, throughput evidence, PLC state, and inspection artifacts — one timestamped API snapshot."),
        ("DIAGNOSE", "Simulation, mock PLC, or unverified calibration makes the authoritative disposition HOLD_REQUIRED."),
        ("TRACEABILITY", "Roll ledger, certificate HMAC, control-audit log. Missing data renders as NO DATA, not 0%."),
    )):
        top = Inches(1.3 + i * 1.35)
        _box(s, Inches(0.45), top, Inches(8.2), Inches(1.22), accent=True)
        add_textbox(s, Inches(0.65), top + Inches(0.12), Inches(7.8), Inches(0.32), [
            {"text": title, "size": 16, "color": ACCENT, "bold": True}
        ])
        add_textbox(s, Inches(0.65), top + Inches(0.48), Inches(7.8), Inches(0.6), [
            {"text": body, "size": 14, "color": INK}
        ])
    _box(s, Inches(8.85), Inches(1.3), Inches(4.0), Inches(5.35), accent=True)
    add_textbox(s, Inches(9.05), Inches(1.45), Inches(3.6), Inches(5.0), [
        {"text": "NOT IN PRODUCTION", "size": 12, "color": RED, "bold": True},
        {"text": "Recipe sliders", "size": 16, "color": INK, "bold": True, "space_after": 4},
        {"text": "Defect injection", "size": 16, "color": INK, "bold": True, "space_after": 4},
        {"text": "7-stage simulator", "size": 16, "color": INK, "bold": True, "space_after": 4},
        {"text": "LIBAD 90s demo", "size": 16, "color": INK, "bold": True, "space_after": 4},
        {"text": "Send offset to PLC", "size": 16, "color": INK, "bold": True, "space_after": 12},
        {"text": "Sandbox-only features stay isolated from the operating surface.", "size": 13, "color": MUTED},
        {"text": "Control = confirm-audit only.", "size": 14, "color": ACCENT, "bold": True},
    ])
    footer(s, 5)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "06", "3:35–4:20")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(12.4), Inches(0.7), [
        {"text": "Intent → command → ACK. HOLD is requested, then confirmed.", "size": 24, "color": INK, "bold": True}
    ])
    steps = [
        ("1", "Inspect", "RGB path + fail-safe deadline"),
        ("2", "PENDING", "SQLite claims (batch, part)"),
        ("3", "PLC", "One command owner + ACK sequence"),
        ("4", "Confirm", "PASS/REJECT only after ACK; HOLD_REQUESTED ≠ HOLD_CONFIRMED"),
        ("5", "Certificate", "HMAC-SHA256 over canonical payload"),
        ("6", "Audit", "Operator confirm phrase is durable"),
    ]
    for i, (number, title, body) in enumerate(steps):
        left = Inches(0.45 + (i % 3) * 4.2)
        top = Inches(1.4 + (i // 3) * 2.5)
        _box(s, left, top, Inches(4.0), Inches(2.25), accent=True)
        add_textbox(s, left + Inches(0.2), top + Inches(0.18), Inches(3.6), Inches(0.35), [
            {"text": number, "size": 14, "color": ACCENT, "bold": True}
        ])
        add_textbox(s, left + Inches(0.2), top + Inches(0.55), Inches(3.6), Inches(0.45), [
            {"text": title, "size": 22, "color": INK, "bold": True}
        ])
        add_textbox(s, left + Inches(0.2), top + Inches(1.15), Inches(3.6), Inches(0.8), [
            {"text": body, "size": 15, "color": MUTED}
        ])
    footer(s, 6)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "07", "4:20–5:10")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(12.4), Inches(0.9), [
        {"text": "Real optical evidence is reproducible — and bounded.", "size": 24, "color": INK, "bold": True}
    ])
    metrics = (
        ("PRECISION", rgb["precision_pct"], f"{rgb['n_test']}-image real optical test split"),
        ("RECALL", rgb["recall_pct"], "measured on configured checkpoint"),
        ("mAP50", rgb["map50_pct"], "measured on configured checkpoint"),
        ("mAP50-95", rgb["map50_95_pct"], "IoU 0.50:0.95"),
    )
    for i, (key, value, note) in enumerate(metrics):
        left = Inches(0.45 + i * 3.15)
        _box(s, left, Inches(1.55), Inches(3.0), Inches(2.15), accent=True)
        add_textbox(s, left + Inches(0.15), Inches(1.68), Inches(2.7), Inches(0.3), [
            {"text": key, "size": 13, "color": MUTED, "bold": True}
        ])
        add_textbox(s, left + Inches(0.15), Inches(2.05), Inches(2.7), Inches(0.55), [
            {"text": value, "size": 28, "color": HOLD if key == "mAP50-95" else INK, "bold": True}
        ])
        add_textbox(s, left + Inches(0.15), Inches(2.7), Inches(2.7), Inches(0.7), [
            {"text": note, "size": 12, "color": MUTED}
        ])
    _box(s, Inches(0.45), Inches(3.95), Inches(12.4), Inches(2.75), accent=True)
    add_textbox(s, Inches(0.7), Inches(4.15), Inches(12.0), Inches(2.4), [
        {"text": "EVIDENCE IDENTITY", "size": 13, "color": ACCENT, "bold": True},
        {"text": (
            f"{rgb['n_test']} test images  ·  seed {rgb['seed']}  ·  "
            f"weights {rgb['weights_prefix']}…  ·  dataset {rgb['dataset_prefix']}…"
        ), "size": 18, "color": INK, "bold": True, "space_after": 10},
        {"text": "Public real optical image-disjoint split. Not roll-disjoint; not factory qualification.", "size": 15, "color": MUTED},
    ])
    footer(s, 7)

    s = prs.slides.add_slide(blank)
    paint_bg(s)
    kicker(s, "08", "5:10–6:00")
    add_textbox(s, Inches(0.45), Inches(0.5), Inches(12.4), Inches(0.45), [
        {"text": "Let the loops run. Then say the close.", "size": 26, "color": INK, "bold": True}
    ])
    _box(s, Inches(0.45), Inches(1.05), Inches(6.15), Inches(4.05))
    add_textbox(s, Inches(0.6), Inches(1.12), Inches(5.85), Inches(0.28), [
        {"text": "RGB LANE  ·  image_1548  ·  HOLD", "size": 12, "color": HOLD, "bold": True}
    ])
    rgb_media = RGB_GIF if RGB_GIF.is_file() else EXTERNAL_DEMO
    # Media frame: y=1.45 → ≤5.05 so CLOSE at y=5.25 never overlaps.
    media_top, media_h = Inches(1.45), Inches(3.55)
    if rgb_media.is_file():
        _add_picture_contain(s, rgb_media, Inches(0.6), media_top, Inches(5.85), media_h)
    _box(s, Inches(6.75), Inches(1.05), Inches(6.15), Inches(4.05))
    add_textbox(s, Inches(6.9), Inches(1.12), Inches(5.85), Inches(0.28), [
        {"text": "LIBAD GATE  ·  fixture PASS / REJECT / REJECT / HOLD", "size": 12, "color": ACCENT, "bold": True}
    ])
    if LIBAD_GIF.is_file():
        _add_picture_contain(s, LIBAD_GIF, Inches(6.9), media_top, Inches(5.85), media_h)
    add_textbox(s, Inches(0.45), Inches(5.25), Inches(12.4), Inches(1.65), [
        {"text": "CLOSE", "size": 12, "color": ACCENT, "bold": True},
        {"text": "The model finds defects. The evidence gate controls when software may authorize a disposition.", "size": 18, "color": INK, "bold": True, "space_after": 8},
        {"text": "Registered title states the target Zero-Trust Edge-Cloud architecture; this artifact validates inspection, evidence, and fail-closed contracts — not completed mTLS/RBAC/OT. GIFs loop checked-in artifacts. Authors' interim is DINOv2 1-seed, not DINOv3.", "size": 13, "color": MUTED},
    ])
    footer(s, 8)

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
